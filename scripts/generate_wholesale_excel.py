import json
import re
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load raw product data
input_path = r'C:/Users/maxim/.gemini/antigravity/brain/856e20be-c83a-42f8-9359-d594b09bf3c3/.system_generated/steps/153/output.txt'
with open(input_path, 'r', encoding='utf-8') as f:
    text_p = json.load(f).get('result', '')

products = json.loads(re.search(r'<untrusted-data-[^>]+>\s*(\[.*?\])\s*</untrusted-data-', text_p, re.DOTALL).group(1))

def round_wholesale(price):
    if price < 10000:
        r = round(price / 100) * 100
        if r % 1000 == 0: r -= 100
        return max(r, 1000)
    elif price < 100000:
        r = round(price / 100) * 100
        rem = r % 1000
        if rem not in (500, 900):
            if rem < 500:
                r = (r // 1000) * 1000 - 100 if rem < 250 else (r // 1000) * 1000 + 500
            else:
                r = (r // 1000) * 1000 + 500 if rem < 750 else (r // 1000) * 1000 + 900
        return r
    else:
        r = round(price / 500) * 500
        if r % 1000 == 0: r -= 100
        return r

cosmetics_brands = {'Medicube', 'SKIN1004', 'Dr. Althea', 'Celimax', 'Anua', 'Numbuzin', 'Lilyeve', 'Tree Hut'}
capilar_brands = {'Shiseido', 'Kérastase', 'Karseell'}
perfume_brands = {'Victoria\'s Secret', 'Lattafa'}

def get_rubro(brand):
    if brand in cosmetics_brands:
        return "Cosmética & Skincare"
    elif brand in capilar_brands:
        return "Cuidado Capilar"
    elif brand in perfume_brands:
        return "Perfumería & Fragancias"
    return "Otros"

calibrated_items = []

for p in products:
    ret = float(p.get('product_retail') or 0)
    landed = float(p.get('landed_cost') or 0)
    orig = float(p.get('origin_cost') or 0)
    brand = p.get('brand', '').strip()
    title = p.get('product_title', '').strip()
    cat = p.get('category_name', 'Cosmética').strip()
    
    if landed <= 0 and orig > 0:
        landed = orig * 1.20 * 1600
    if landed <= 0:
        landed = ret * 0.55
        
    # Check broken retail reference
    adjusted_retail = ret
    if ret < landed * 1.15:
        adjusted_retail = round(landed * 1.60 / 100) * 100
        
    mya_floor = landed * 1.18 # MYA markup floor 18%
    reseller_target = adjusted_retail * 0.74 # 26% discount off retail
    
    if reseller_target >= mya_floor:
        target_whs = reseller_target
    else:
        if adjusted_retail > mya_floor * 1.08:
            target_whs = mya_floor
        else:
            target_whs = (adjusted_retail + landed) / 2
            
    final_whs = round_wholesale(target_whs)
    if final_whs >= adjusted_retail:
        final_whs = round_wholesale(adjusted_retail * 0.82)
    if final_whs < landed:
        final_whs = round_wholesale(landed * 1.10)
        
    # Minimum quantity rule
    min_qty = int(p.get('product_min_qty') or 4)
    if 'mascarilla' in title.lower() and 'pack' not in title.lower() and final_whs < 10000:
        min_qty = 10
    elif final_whs > 100000:
        min_qty = 3
    elif brand == 'Lattafa':
        min_qty = 3
    else:
        min_qty = 4
        
    calibrated_items.append({
        'id': p['product_id'],
        'sku': p.get('product_sku') or f"MYA-{p['product_id'][:8].upper()}",
        'brand': brand,
        'title': title,
        'category_name': cat,
        'retail_price': adjusted_retail,
        'wholesale_price': final_whs,
        'min_qty': min_qty,
        'rubro': get_rubro(brand)
    })

cosmetics_items = sorted([item for item in calibrated_items if item['rubro'] == "Cosmética & Skincare"], key=lambda x: (x['brand'], x['title']))
capilar_items = sorted([item for item in calibrated_items if item['rubro'] == "Cuidado Capilar"], key=lambda x: (x['brand'], x['title']))
perfume_items = sorted([item for item in calibrated_items if item['rubro'] == "Perfumería & Fragancias"], key=lambda x: (x['brand'], x['title']))
all_items = sorted(calibrated_items, key=lambda x: (x['rubro'], x['brand'], x['title']))

# Create Workbook
wb = openpyxl.Workbook()
wb.remove(wb.active)

# Color Scheme: Midnight Navy & Emerald Luxury
NAVY_PRIMARY = "0F172A"       # Deep Slate 900
NAVY_HEADER = "1E293B"        # Slate 800
ACCENT_CYAN = "0284C7"        # Sky 600
BG_KPI = "F8FAFC"             # Slate 50
BORDER_LIGHT = "E2E8F0"       # Slate 200
BORDER_DARK = "334155"        # Slate 700
ROW_ALT = "F8FAFC"            # Slate 50
WHITE = "FFFFFF"
EMERALD_BG = "ECFDF5"         # Mint/Emerald 50
EMERALD_TEXT = "065F46"       # Emerald 800
GRAY_TEXT = "64748B"          # Slate 500

# Fonts
font_brand_title = Font(name="Segoe UI", size=15, bold=True, color="0F172A")
font_section_title = Font(name="Segoe UI", size=10, bold=True, color="0284C7")
font_kpi_text = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")

font_col_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
fill_col_header = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")

font_regular = Font(name="Segoe UI", size=9.5, color="1E293B")
font_sku = Font(name="Consolas", size=9, color="475569")
font_brand = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
font_retail = Font(name="Segoe UI", size=9.5, color=GRAY_TEXT)

font_wholesale = Font(name="Segoe UI", size=10.5, bold=True, color=EMERALD_TEXT)
fill_wholesale = PatternFill(start_color=EMERALD_BG, end_color=EMERALD_BG, fill_type="solid")

fill_alt = PatternFill(start_color=ROW_ALT, end_color=ROW_ALT, fill_type="solid")
fill_white = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
fill_kpi = PatternFill(start_color=BG_KPI, end_color=BG_KPI, fill_type="solid")

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

thin_border = Border(
    left=Side(style='thin', color=BORDER_LIGHT),
    right=Side(style='thin', color=BORDER_LIGHT),
    top=Side(style='thin', color=BORDER_LIGHT),
    bottom=Side(style='thin', color=BORDER_LIGHT)
)

kpi_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

header_border = Border(
    left=Side(style='thin', color=BORDER_DARK),
    right=Side(style='thin', color=BORDER_DARK),
    top=Side(style='medium', color=NAVY_PRIMARY),
    bottom=Side(style='medium', color=NAVY_PRIMARY)
)

# -------------------------------------------------------------
# 1. COVER SHEET: Inicio, Condiciones & Directorio
# -------------------------------------------------------------
ws_cover = wb.create_sheet(title="Inicio & Condiciones")
ws_cover.views.sheetView[0].showGridLines = True

ws_cover.column_dimensions["A"].width = 5
ws_cover.column_dimensions["B"].width = 30
ws_cover.column_dimensions["C"].width = 82

ws_cover.row_dimensions[2].height = 32
ws_cover["B2"] = "MYA IMPORTACIONES · CATÁLOGO MAYORISTA OFICIAL"
ws_cover["B2"].font = Font(name="Segoe UI", size=17, bold=True, color="0F172A")

ws_cover.row_dimensions[3].height = 20
ws_cover["B3"] = "Cosmética Coreana (K-Beauty), Cuidado Capilar Profesional y Perfumería de Autor"
ws_cover["B3"].font = Font(name="Segoe UI", size=11, bold=True, color=ACCENT_CYAN)

cover_info = [
    ("Mínimo de Compra", "Mínimo de $100.000 ARS en productos O 5 unidades en total. Podés combinar libremente cualquier marca, categoría y rubro comercial."),
    ("Logística & Despacho", "Envíos asegurados a todo el país por Correo Argentino (modalidad oficial Paq.ar a domicilio o a sucursal). Te enviamos el código de seguimiento en vivo por WhatsApp al momento del despacho."),
    ("Punto de Retiro en Tandil", "Si te encontrás en Tandil o zona de influencia, podés coordinar el retiro sin costo de envío directamente en nuestra sede central."),
    ("Embalaje de Seguridad", "Toda la cosmética y perfumería se acondiciona con protección multicapa reforzada sin costo extra para garantizar que llegue intacta."),
    ("Beneficio de Pago Directo", "3% de descuento adicional abonando por Transferencia Bancaria directa o depósito en cuenta oficial MYA."),
    ("Garantía Oficial MYA", "Garantía directa de 30 días sobre cualquier falla de origen o rotura en tránsito."),
    ("Atención y Pedidos", "Atención humana personalizada por WhatsApp con Agustina y Máximo para armar pedidos, consultar stock y coordinar despachos inmediatos.")
]

r_idx = 5
for title, desc in cover_info:
    ws_cover.row_dimensions[r_idx].height = 26
    c_title = ws_cover.cell(row=r_idx, column=2, value=title)
    c_title.font = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
    c_title.fill = fill_kpi
    c_title.border = kpi_border
    c_title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    
    c_desc = ws_cover.cell(row=r_idx, column=3, value=desc)
    c_desc.font = Font(name="Segoe UI", size=10, color="334155")
    c_desc.border = thin_border
    c_desc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    r_idx += 1

# Directory block
r_idx += 2
ws_cover.row_dimensions[r_idx].height = 26
ws_cover.cell(row=r_idx, column=2, value="SECCIONES DEL CATÁLOGO").font = Font(name="Segoe UI", size=12, bold=True, color="0F172A")
r_idx += 1

sections = [
    ("🇰🇷 Pestaña: Cosmética & Skincare", f"{len(cosmetics_items)} artículos (Medicube, SKIN1004, Celimax, Anua, Dr. Althea, Numbuzin, Lilyeve, Tree Hut)"),
    ("💆 Pestaña: Cuidado Capilar", f"{len(capilar_items)} artículos (Shiseido Tsubaki/Fino, Kérastase profesional, Karseell Maca)"),
    ("🌸 Pestaña: Perfumería & Fragancias", f"{len(perfume_items)} artículos (Victoria's Secret Body Mists/Lociones, Lattafa Perfumes Árabes)"),
    ("🌟 Pestaña: Catálogo General Completo", f"{len(all_items)} artículos unificados con clasificación por Rubro Comercial")
]

for s_name, s_detail in sections:
    ws_cover.row_dimensions[r_idx].height = 24
    c_s = ws_cover.cell(row=r_idx, column=2, value=s_name)
    c_s.font = Font(name="Segoe UI", size=10.5, bold=True, color="1E3A8A")
    c_s.border = thin_border
    
    c_d = ws_cover.cell(row=r_idx, column=3, value=s_detail)
    c_d.font = Font(name="Segoe UI", size=10, color="475569")
    c_d.border = thin_border
    r_idx += 1

# Function to build product sheets
def build_product_sheet(ws, sheet_title, items, is_general=False):
    ws.views.sheetView[0].showGridLines = True
    
    # Row Heights
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 30
    ws.row_dimensions[4].height = 10
    ws.row_dimensions[5].height = 28
    
    # Title Banner
    ws["A1"] = "MYA IMPORTACIONES · LISTADO OFICIAL DE PRECIOS MAYORISTAS"
    ws["A1"].font = font_brand_title
    
    ws["A2"] = f"{sheet_title.upper()} · VIGENCIA OCTUBRE 2026 · TANDIL, ARGENTINA"
    ws["A2"].font = font_section_title
    
    # KPI Badges
    ws["A3"] = "📦 MÍNIMO GENERAL:\n$100.000 o 5 u. combinadas"
    ws["A3"].font = font_kpi_text
    ws["A3"].fill = fill_kpi
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["A3"].border = kpi_border
    
    ws["B3"] = "🚚 LOGÍSTICA:\nCorreo Argentino (Paq.ar) / Retiro Tandil"
    ws["B3"].font = font_kpi_text
    ws["B3"].fill = fill_kpi
    ws["B3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["B3"].border = kpi_border
    
    badge_c = "D3" if is_general else "C3"
    ws[badge_c] = "💳 FORMA DE PAGO:\nTransferencia bancaria (-3% OFF)"
    ws[badge_c].font = font_kpi_text
    ws[badge_c].fill = fill_kpi
    ws[badge_c].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws[badge_c].border = kpi_border
    
    # Headers
    if is_general:
        headers = [
            "Código / SKU",
            "Marca",
            "Rubro Comercial",
            "Producto / Presentación",
            "Categoría",
            "Mínimo Mayorista",
            "Precio Minorista (Ref. PVP)",
            "Precio Mayorista MYA"
        ]
    else:
        headers = [
            "Código / SKU",
            "Marca",
            "Producto / Presentación",
            "Categoría",
            "Mínimo Mayorista",
            "Precio Minorista (Ref. PVP)",
            "Precio Mayorista MYA"
        ]
        
    start_row = 5
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c_idx, value=h)
        cell.font = font_col_header
        cell.fill = fill_col_header
        cell.alignment = align_center
        cell.border = header_border
        
    ws.freeze_panes = f"A{start_row + 1}"
    
    curr = start_row + 1
    for i, itm in enumerate(items):
        ws.row_dimensions[curr].height = 21
        row_fill = fill_alt if i % 2 == 1 else fill_white
        
        col = 1
        # SKU
        c = ws.cell(row=curr, column=col, value=itm['sku'])
        c.font = font_sku; c.fill = row_fill; c.alignment = align_center; c.border = thin_border
        col += 1
        
        # Marca
        c = ws.cell(row=curr, column=col, value=itm['brand'])
        c.font = font_brand; c.fill = row_fill; c.alignment = align_left; c.border = thin_border
        col += 1
        
        # Rubro (if general)
        if is_general:
            c = ws.cell(row=curr, column=col, value=itm['rubro'])
            c.font = font_regular; c.fill = row_fill; c.alignment = align_left; c.border = thin_border
            col += 1
            
        # Title
        c = ws.cell(row=curr, column=col, value=itm['title'])
        c.font = font_regular; c.fill = row_fill; c.alignment = align_left; c.border = thin_border
        col += 1
        
        # Category
        c = ws.cell(row=curr, column=col, value=itm['category_name'])
        c.font = font_regular; c.fill = row_fill; c.alignment = align_left; c.border = thin_border
        col += 1
        
        # Min Qty
        c = ws.cell(row=curr, column=col, value=f"{itm['min_qty']} u.")
        c.font = font_regular; c.fill = row_fill; c.alignment = align_center; c.border = thin_border
        col += 1
        
        # Retail Price
        c = ws.cell(row=curr, column=col, value=itm['retail_price'])
        c.font = font_retail; c.fill = row_fill; c.alignment = align_right; c.border = thin_border
        c.number_format = '"$ "#,##0'
        col += 1
        
        # Wholesale Price
        c = ws.cell(row=curr, column=col, value=itm['wholesale_price'])
        c.font = font_wholesale; c.fill = fill_wholesale; c.alignment = align_right; c.border = thin_border
        c.number_format = '"$ "#,##0'
        
        curr += 1
        
    last_row = curr - 1
    last_col = len(headers)
    last_col_letter = get_column_letter(last_col)
    
    # AutoFilter
    ws.auto_filter.ref = f"A{start_row}:{last_col_letter}{last_row}"
    
    # Column auto-sizing
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row < start_row:
                continue
            v_str = str(cell.value or '')
            if cell.number_format and '$' in cell.number_format:
                v_str = "$ " + v_str + "    "
            max_len = max(max_len, len(v_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
        
    # Extra width for product title column
    title_col_letter = "D" if is_general else "C"
    ws.column_dimensions[title_col_letter].width = 50

# -------------------------------------------------------------
# Sheet 2: Cosmética & Skincare
# -------------------------------------------------------------
ws_cosm = wb.create_sheet(title="Cosmética & Skincare")
build_product_sheet(ws_cosm, "Cosmética Coreana & Skincare (166 Artículos)", cosmetics_items, is_general=False)

# -------------------------------------------------------------
# Sheet 3: Cuidado Capilar
# -------------------------------------------------------------
ws_cap = wb.create_sheet(title="Cuidado Capilar")
build_product_sheet(ws_cap, "Cuidado Capilar & Tratamientos (48 Artículos)", capilar_items, is_general=False)

# -------------------------------------------------------------
# Sheet 4: Perfumería & Fragancias
# -------------------------------------------------------------
ws_perf = wb.create_sheet(title="Perfumería & Fragancias")
build_product_sheet(ws_perf, "Perfumería Árabe & Body Mists (40 Artículos)", perfume_items, is_general=False)

# -------------------------------------------------------------
# Sheet 5: Catálogo Completo General
# -------------------------------------------------------------
ws_gen = wb.create_sheet(title="Catálogo General Completo")
build_product_sheet(ws_gen, "Catálogo Mayorista General (254 Artículos)", all_items, is_general=True)

# Save files safely (in case user has one file open in Excel)
xlsx_destinations = [
    r"c:\Users\maxim\OneDrive\Escritorio\MYAimportaciones\MYA_Catalogo_Mayorista_Cosmetica_Premium.xlsx",
    r"c:\Users\maxim\OneDrive\Escritorio\MYAimportaciones\Tienda WEB\public\MYA_Precios_Mayoristas_Cosmetica.xlsx",
    r"c:\Users\maxim\OneDrive\Escritorio\MYAimportaciones\MYA_Precios_Mayoristas_Cosmetica.xlsx"
]

for path in xlsx_destinations:
    try:
        wb.save(path)
        print(f"Premium Excel saved successfully to: {path}")
    except PermissionError:
        print(f"File {path} is currently open in Excel, skipped overwrite.")

# Save clean CSV
csv_path = r"c:\Users\maxim\OneDrive\Escritorio\MYAimportaciones\MYA_Precios_Mayoristas_Cosmetica.csv"
import csv
try:
    with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f, delimiter=';')
        writer.writerow(["SKU", "Marca", "Rubro", "Producto / Presentacion", "Categoria", "Minimo Mayorista", "Precio Minorista (Ref PVP)", "Precio Mayorista MYA"])
        for itm in all_items:
            writer.writerow([
                itm['sku'],
                itm['brand'],
                itm['rubro'],
                itm['title'],
                itm['category_name'],
                f"{itm['min_qty']} u.",
                int(itm['retail_price']),
                int(itm['wholesale_price'])
            ])
    print(f"Premium CSV saved successfully to: {csv_path}")
except PermissionError:
    print(f"CSV {csv_path} is currently open, skipped overwrite.")

