import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
OUT=Path(__file__).resolve().parents[1]/'docs/electronics-atacado-2026-10-08'
rows=json.loads((OUT/'priced.json').read_text(encoding='utf8'))
book=Workbook();s=book.active;s.title='Revisión privada'
headers=['SKU','Marca','Modelo y variante','Compra USD','Compra ARS','Reventa USD','Reventa ARS','Diferencia antes de gastos ARS','Recargo sobre compra %','Estado','Pendientes','Comparable ARS','Techo competitivo ARS','Fuente comparable','Documento','Fila']
s.append(headers)
for r in rows:
 s.append([r['sku'],r['brand'],r['title'],r['purchase_usd'],r['purchase_ars'],r['resale_usd'],r['retail_price'],r['gross_before_expenses'],r['purchase_markup_percent'],'Apto' if r['eligible'] else 'Borrador',' | '.join(r['pending']),r['reference_ars'],r['competitive_ceiling_ars'],(r['reference'] or {}).get('url',''),r['file'],r['row']])
 for col in [4,5,6,7,8,12,13]:s.cell(s.max_row,col).number_format='#,##0.00'
 s.cell(s.max_row,10).fill=PatternFill('solid',fgColor='DBF4DF' if r['eligible'] else 'FFF0CC')
for cell in s[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='18324F');cell.alignment=Alignment(wrap_text=True)
s.freeze_panes='D2';s.auto_filter.ref=s.dimensions;s.row_dimensions[1].height=44
for col in ['A','B','D','E','F','G','H','I','J','L','M','P']:s.column_dimensions[col].width=20
for col,width in [('C',65),('K',65),('N',55),('O',38)]:s.column_dimensions[col].width=width
notes=book.create_sheet('Criterio');notes.append(['Nota','Detalle'])
for row in [('Cambio USD/ARS','1550 · DolarAPI blue venta 08/10/2026'),('Criterio delegado','Rentabilidad y precio competitivo: 3% bajo techo competitivo, redondeado hacia abajo a ARS100.'),('Piso inicial','15% sobre compra antes de gastos; no garantiza rentabilidad neta.'),('Gastos pendientes','Flete, internación, comisiones, embalaje y otros gastos desconocidos. Cero obligatorio en la base no significa gratis.'),('Borradores sin referencia','Compra +25% es sólo una propuesta provisional privada; no comparable ni precio certificado para activar.'),('Stock','Sin stock físico propio confirmado. Oferta a pedido según lista, disponibilidad de la unidad a confirmar.')]:notes.append(row)
notes.column_dimensions['A'].width=28;notes.column_dimensions['B'].width=115
book.save(OUT/'revision-privada.xlsx')
print('Private review workbook: 137 variants, costs and pending reasons.')
